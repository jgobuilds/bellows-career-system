#!/usr/bin/env python3
"""
cover_builder.py — render a cover letter .docx from a JSON spec.

Companion to resume_builder.py. Content lives in the spec; the clean layout and
the placeholder check live here. Keep the letter honest and specific per
resume-style-rules.md §10-11 (portability test on the opener, candor on any
gap) — this builder only renders; the judgment is in the paragraphs you pass it.

USAGE
  python engine/cover_builder.py <spec.json> <out.docx>

SPEC SHAPE:
  {
    "name": "Firstname Lastname, M.S.",
    "contact": "City, ST  |  email  |  phone  |  linkedin",
    "subject": "Re: Head of Data & AI",
    "paragraphs": ["opening (must pass the portability test)", "body", "..."],
    "signoff": "Firstname Lastname"
  }
"""

import json
import sys

import docx
from docx.shared import Inches, Pt

from docx_common import BODY_FONT as FONT
from docx_common import BRAND_BLUE, run, scan_placeholders


def build_cover(spec, out_path):
    warns = [f"unresolved placeholder: {m!r}" for m in sorted(scan_placeholders(spec))]

    d = docx.Document()
    for s in d.sections:
        s.top_margin = s.bottom_margin = Inches(0.7)
        s.left_margin = s.right_margin = Inches(0.8)
    st = d.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.08

    def p(text="", bold=False, size=10.5, after=8, color=None):
        par = d.add_paragraph()
        par.paragraph_format.space_after = Pt(after)
        if text:
            run(par, text, bold=bold, size=size, color=color, font=FONT)
        return par

    p(spec["name"], bold=True, size=16, after=1, color=BRAND_BLUE)
    p(spec["contact"], size=9.5, after=14)
    if spec.get("subject"):
        p(spec["subject"], bold=True, after=10)
    for para in spec["paragraphs"]:
        p(para)
    p(spec.get("closing", "Thank you for your time and consideration."), after=12)
    p(spec["signoff"], bold=True)

    d.save(out_path)
    return warns


def main():
    if len(sys.argv) != 3:
        sys.exit("usage: python cover_builder.py <spec.json> <out.docx>")
    spec = json.load(open(sys.argv[1], encoding="utf-8"))
    warns = build_cover(spec, sys.argv[2])
    print(f"built {sys.argv[2]}")
    for w in warns:
        print("   ⚠", w)


if __name__ == "__main__":
    main()
